from __future__ import annotations

import json
from pathlib import Path
from typing import Callable

from ..chunking import ParsedChunk
from ..errors import ErrorCode, RagError

PARSER_VERSION = "1.0.0"

TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv",
    ".json", ".xml", ".yaml", ".yml",
}
CODE_EXTS = {
    ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".go", ".rs", ".sql",
    ".c", ".h", ".cpp", ".hpp", ".cs", ".rb", ".php", ".sh", ".ps1", ".bat", ".toml", ".ini", ".cfg",
}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
PPTX_EXTS = {".pptx"}
XLSX_EXTS = {".xlsx", ".xlsm"}
HTML_EXTS = {".html", ".htm"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
AUDIO_EXTS = {".wav", ".mp3", ".m4a", ".aac", ".flac", ".ogg"}
VIDEO_EXTS = {".mp4", ".mov", ".mkv", ".webm", ".avi"}

TEXTUAL_EXTS = TEXT_EXTS | CODE_EXTS
SUPPORTED_EXTS = (
    TEXTUAL_EXTS | PDF_EXTS | DOCX_EXTS | PPTX_EXTS | XLSX_EXTS
    | HTML_EXTS | IMAGE_EXTS | AUDIO_EXTS | VIDEO_EXTS
)


def extension_of(path: Path) -> str:
    return path.suffix.lower()


def is_supported(path: Path) -> bool:
    return extension_of(path) in SUPPORTED_EXTS


def file_type_of(path: Path) -> str:
    return extension_of(path).lstrip(".")


def require(cond: bool, code: str, msg: str, suggestion: str = "") -> None:
    if not cond:
        raise RagError(code, msg, retryable=False, suggestion=suggestion)


def parse_text_file(path: Path, max_chars: int, overlap: int) -> list[ParsedChunk]:
    from charset_normalizer import from_path

    try:
        text = str(from_path(path).best())
    except Exception:
        text = path.read_text(encoding="utf-8", errors="replace")
    if not text.strip():
        return []
    from ..chunking import chunk_text

    track = extension_of(path) in CODE_EXTS
    return chunk_text(text, max_chars=max_chars, overlap=overlap, track_lines=track)


def parse_html_file(path: Path, max_chars: int, overlap: int) -> list[ParsedChunk]:
    try:
        from bs4 import BeautifulSoup
        raw = path.read_bytes()
        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        title = soup.title.get_text(" ", strip=True) if soup.title else ""
        text = soup.get_text("\n", strip=True)
        if title:
            text = title + "\n\n" + text
    except ImportError as e:
        raise RagError(
            ErrorCode.PARSER_DEPENDENCY_MISSING,
            "beautifulsoup4 is required for HTML parsing",
            suggestion="pip install beautifulsoup4",
        ) from e
    from ..chunking import chunk_text

    return chunk_text(text, max_chars=max_chars, overlap=overlap)


def parse_pdf_file(path: Path, max_chars: int, overlap: int, ctx) -> list[ParsedChunk]:
    """Text-layer PDF via pypdf, per page. Scanned pages require vision (PRD 6.2)."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    chunks: list[ParsedChunk] = []
    empty_pages: list[int] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        text = text.strip()
        if not text:
            empty_pages.append(i)
            continue
        from ..chunking import chunk_text

        for c in chunk_text(text, max_chars=max_chars, overlap=overlap):
            c.page_number = i
            chunks.append(c)
    if empty_pages and ctx is not None:
        chunks.extend(_pdf_pages_via_vision(path, ctx, empty_pages))
    if empty_pages and ctx is None:
        from ..errors import RagError

        raise RagError(
            ErrorCode.VISION_PROVIDER_NOT_CONFIGURED,
            f"PDF has no text layer on pages {empty_pages[:5]}{'...' if len(empty_pages) > 5 else ''} "
            "(scanned?) and vision/OCR is not configured",
            suggestion="Configure [vision] provider, or use a text-layer PDF.",
            details={"empty_pages": empty_pages},
        )
    return chunks


def _pdf_pages_via_vision(path: Path, ctx, pages: list[int]) -> list[ParsedChunk]:
    """Render pages to PNG (requires pymupdf) and caption via vision provider."""
    try:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
    except ImportError as e:
        raise RagError(
            ErrorCode.PARSER_DEPENDENCY_MISSING,
            "pymupdf is required to render scanned PDF pages for vision OCR",
            suggestion="pip install pymupdf, and configure [vision].",
        ) from e
    vision = ctx.vision_provider()
    chunks: list[ParsedChunk] = []
    doc = fitz.open(str(path))
    import tempfile, os

    for pno in pages[:20]:
        page = doc[pno - 1]
        pix = page.get_pixmap(dpi=150)
        tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
        tmp.close()
        pix.save(tmp.name)
        try:
            caption = vision.describe_image(tmp.name, task="ocr")
        finally:
            os.unlink(tmp.name)
        from ..chunking import ParsedChunk

        chunks.append(
            ParsedChunk(
                content=caption,
                content_type="image_ocr",
                origin_kind="model_generated",
                page_number=pno,
                metadata={"model": vision.model, "via": "vision_api_ocr"},
            )
        )
    doc.close()
    return chunks


def parse_docx_file(path: Path, max_chars: int, overlap: int) -> list[ParsedChunk]:
    try:
        import docx
    except ImportError as e:
        raise RagError(ErrorCode.PARSER_DEPENDENCY_MISSING, "python-docx required") from e
    document = docx.Document(str(path))
    from ..chunking import chunk_text, ParsedChunk

    chunks: list[ParsedChunk] = []
    buf: list[str] = []

    def flush():
        if buf:
            for c in chunk_text("\n\n".join(buf), max_chars=max_chars, overlap=overlap):
                chunks.append(c)
            buf.clear()

    for para in document.paragraphs:
        t = para.text.strip()
        if para.style.name.startswith("Heading"):
            flush()
            if t:
                c = ParsedChunk(content=f"# {t}", content_type="text", origin_kind="extracted_text")
                chunks.append(c)
        elif t:
            buf.append(t)
    flush()
    for t_idx, table in enumerate(document.tables):
        rows = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            chunks.append(ParsedChunk(
                content="\n".join(rows), content_type="table", origin_kind="extracted_text",
                metadata={"table_index": t_idx},
            ))
    return chunks


def parse_pptx_file(path: Path, max_chars: int, overlap: int) -> list[ParsedChunk]:
    try:
        from pptx import Presentation
    except ImportError as e:
        raise RagError(ErrorCode.PARSER_DEPENDENCY_MISSING, "python-pptx required") from e
    prs = Presentation(str(path))
    from ..chunking import chunk_text

    chunks: list[ParsedChunk] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    parts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = []
                for row in shape.table.rows:
                    rows.append(" | ".join(c.text.strip() for c in row.cells))
                if rows:
                    chunks.append(ParsedChunk(
                        content="\n".join(rows), content_type="table",
                        slide_number=i, origin_kind="extracted_text",
                    ))
        if parts:
            for c in chunk_text("\n".join(parts), max_chars=max_chars, overlap=overlap):
                c.slide_number = i
                chunks.append(c)
    return chunks


def parse_xlsx_file(path: Path, max_chars: int, overlap: int) -> list[ParsedChunk]:
    try:
        import openpyxl
    except ImportError as e:
        raise RagError(ErrorCode.PARSER_DEPENDENCY_MISSING, "openpyxl required") from e
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    from ..chunking import ParsedChunk

    chunks: list[ParsedChunk] = []
    for ws in wb.worksheets:
        lines: list[str] = []
        header = None
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = ["" if v is None else str(v) for v in row]
            if not any(c.strip() for c in cells):
                continue
            line = " | ".join(cells)
            if r_idx == 1:
                header = line
                lines.append(line)
            else:
                lines.append(line)
            if len("\n".join(lines)) >= max_chars:
                chunks.append(ParsedChunk(
                    content="\n".join(lines), content_type="table", sheet_name=ws.title,
                    origin_kind="extracted_text",
                    metadata={"row_start": 1, "row_end": r_idx},
                ))
                lines = ([header] if header else [])
        if lines:
            chunks.append(ParsedChunk(
                content="\n".join(lines), content_type="table", sheet_name=ws.title,
                origin_kind="extracted_text",
            ))
    wb.close()
    return chunks


def parse_image_file(path: Path, ctx) -> list[ParsedChunk]:
    require(ctx is not None and ctx.vision_configured(),
            ErrorCode.VISION_PROVIDER_NOT_CONFIGURED,
            "Image indexing requires a configured vision provider",
            "Configure [vision] provider/base_url/model and its API key env var.")
    vision = ctx.vision_provider()
    caption = vision.describe_image(path, task="caption")
    return [ParsedChunk(
        content=caption, content_type="image_caption", origin_kind="model_generated",
        metadata={"model": vision.model},
    )]


def parse_audio_file(path: Path, ctx) -> list[ParsedChunk]:
    require(ctx is not None and ctx.asr_configured(),
            ErrorCode.ASR_PROVIDER_NOT_CONFIGURED,
            "Audio indexing requires a configured ASR provider",
            "Configure [asr] provider/base_url/model and its API key env var.")
    from ..chunking import ParsedChunk

    asr = ctx.asr_provider()
    result = asr.transcribe(str(path))
    chunks: list[ParsedChunk] = []
    if result["segments"]:
        for seg in result["segments"]:
            text = (seg.get("text") or "").strip()
            if not text:
                continue
            chunks.append(ParsedChunk(
                content=text, content_type="audio_transcript", origin_kind="extracted_text",
                start_ms=int((seg.get("start") or 0) * 1000),
                end_ms=int((seg.get("end") or 0) * 1000),
                metadata={"timestamp_precision": result["timestamp_precision"], "model": asr.model},
            ))
    elif result["text"].strip():
        chunks.append(ParsedChunk(
            content=result["text"].strip(), content_type="audio_transcript",
            origin_kind="extracted_text",
            metadata={"timestamp_precision": "unavailable", "model": asr.model},
        ))
    return chunks


def parse_video_file(path: Path, ctx) -> list[ParsedChunk]:
    import shutil

    ffmpeg = ctx.ffmpeg_path() if ctx else shutil.which("ffmpeg")
    if not ffmpeg:
        raise RagError(
            ErrorCode.PARSER_DEPENDENCY_MISSING,
            "ffmpeg not found; video indexing requires FFmpeg",
            suggestion="Install FFmpeg and add it to PATH, or set [video].ffmpeg_path.",
        )
    from ..chunking import ParsedChunk

    chunks: list[ParsedChunk] = []
    # 1) audio track -> ASR
    import tempfile, os, subprocess

    have_asr = ctx is not None and ctx.asr_configured()
    audio_warnings: list[str] = []
    tmp_wav = None
    if have_asr:
        fd, tmp_wav = tempfile.mkstemp(suffix=".wav")
        os.close(fd)
        try:
            subprocess.run(
                [ffmpeg, "-y", "-i", str(path), "-vn", "-acodec", "pcm_s16le",
                 "-ar", "16000", "-ac", "1", tmp_wav],
                capture_output=True, timeout=3600,
            )
            if os.path.getsize(tmp_wav) > 1024:
                asr = ctx.asr_provider()
                result = asr.transcribe(tmp_wav)
                for seg in result["segments"] or ([{"start": 0, "end": None, "text": result["text"]}] if result["text"].strip() else []):
                    text = (seg.get("text") or "").strip()
                    if text:
                        chunks.append(ParsedChunk(
                            content=text, content_type="video_transcript", origin_kind="extracted_text",
                            start_ms=int((seg.get("start") or 0) * 1000),
                            end_ms=int(seg["end"] * 1000) if seg.get("end") is not None else None,
                            metadata={"model": asr.model, "timestamp_precision": result["timestamp_precision"]},
                        ))
        finally:
            if tmp_wav and os.path.exists(tmp_wav):
                os.unlink(tmp_wav)
    elif ctx is not None and ctx.cfg.asr.enabled:
        # graceful degradation: index frames now, record honestly that speech was not transcribed
        audio_warnings.append(
            "ASR_PROVIDER_NOT_CONFIGURED: audio track present but not transcribed; configure [asr]."
        )

    # 2) keyframes -> vision captions (if configured)
    if ctx is not None and ctx.vision_configured() and ctx.cfg.video.enabled:
        import json as _json

        fps = max(1.0 / 60, min(6 / 60, ctx.cfg.video.max_frames_per_minute / 60))
        fd, tmpdir = tempfile.mkstemp(suffix=".frames")
        os.close(fd)
        os.unlink(tmpdir)
        os.makedirs(tmpdir, exist_ok=True)
        try:
            subprocess.run(
                [ffmpeg, "-y", "-i", str(path), "-vf", f"fps={fps}", "-frames:v",
                 str(ctx.cfg.video.max_frames_per_video), os.path.join(tmpdir, "f%05d.png")],
                capture_output=True, timeout=3600,
            )
            frames = sorted(os.listdir(tmpdir))
            vision = ctx.vision_provider()
            for fname in frames[: ctx.cfg.video.max_frames_per_video]:
                fpath = os.path.join(tmpdir, fname)
                seq = int(fname[1:6])
                ts_ms = int(seq / fps * 1000)
                try:
                    caption = vision.describe_image(fpath)
                except RagError:
                    continue
                chunks.append(ParsedChunk(
                    content=caption, content_type="video_frame_caption",
                    origin_kind="model_generated", start_ms=ts_ms, end_ms=None,
                    metadata={"model": vision.model, **({"audio_warnings": audio_warnings} if audio_warnings else {})},
                ))
        finally:
            import shutil as _sh

            _sh.rmtree(tmpdir, ignore_errors=True)
    if not chunks:
        raise RagError(
            ErrorCode.PARSER_FAILED,
            "Video produced no indexable content (no decodable frames or audio)",
            suggestion="Check the video file is valid; configure [asr]/[vision].",
        )
    return chunks


def parse_file(path: Path, ctx, *, max_chars: int, overlap: int) -> tuple[str, list[ParsedChunk]]:
    """Returns (parser_name, chunks). Raises RagError on failure."""
    ext = extension_of(path)
    if ext in TEXT_EXTS:
        return "text", parse_text_file(path, max_chars, overlap)
    if ext in CODE_EXTS:
        return "code", parse_text_file(path, max_chars, overlap)
    if ext in HTML_EXTS:
        return "html", parse_html_file(path, max_chars, overlap)
    if ext in PDF_EXTS:
        return "pdf", parse_pdf_file(path, max_chars, overlap, ctx)
    if ext in DOCX_EXTS:
        return "docx", parse_docx_file(path, max_chars, overlap)
    if ext in PPTX_EXTS:
        return "pptx", parse_pptx_file(path, max_chars, overlap)
    if ext in XLSX_EXTS:
        return "xlsx", parse_xlsx_file(path, max_chars, overlap)
    if ext in IMAGE_EXTS:
        return "image", parse_image_file(path, ctx)
    if ext in AUDIO_EXTS:
        return "audio", parse_audio_file(path, ctx)
    if ext in VIDEO_EXTS:
        return "video", parse_video_file(path, ctx)
    raise RagError(
        ErrorCode.UNSUPPORTED_FILE_TYPE,
        f"Unsupported file type: {ext or '(no extension)'}",
        details={"extension": ext},
    )
