from __future__ import annotations

import pytest

from app.parsers import parse_docx_file, parse_file, parse_html_file, parse_pptx_file, parse_xlsx_file


def test_markdown_and_text(tmp_path):
    p = tmp_path / "note.md"
    p.write_text("# Title\n\nHello **world** paragraph.\n\nSecond paragraph.\n", encoding="utf-8")
    parser, chunks = parse_file(p, None, max_chars=500, overlap=50)
    assert parser == "text"
    assert any("Hello" in c.content for c in chunks)


def test_csv_json(tmp_path):
    p = tmp_path / "data.json"
    p.write_text('{"key": "特种阀门压力值", "n": 42}', encoding="utf-8")
    parser, chunks = parse_file(p, None, max_chars=500, overlap=50)
    assert parser == "text"
    assert any("特种阀门" in c.content for c in chunks)


def test_unsupported(tmp_path):
    p = tmp_path / "prog.exe"
    p.write_bytes(b"MZ...")
    from app.errors import RagError

    with pytest.raises(RagError) as ei:
        parse_file(p, None, max_chars=500, overlap=50)
    assert ei.value.code == "UNSUPPORTED_FILE_TYPE"


def test_docx(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Pump Overview", level=1)
    d.add_paragraph("The valve opens and liquid level rises.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "item"
    t.cell(0, 1).text = "value"
    t.cell(1, 0).text = "valve"
    t.cell(1, 1).text = "open"
    p = tmp_path / "doc.docx"
    d.save(str(p))
    chunks = parse_docx_file(p, max_chars=800, overlap=50)
    assert any("valve opens" in c.content for c in chunks)
    assert any(c.content_type == "table" for c in chunks)


def test_xlsx(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sensors"
    ws.append(["sensor", "reading"])
    ws.append(["level_1", 3.42])
    p = tmp_path / "book.xlsx"
    wb.save(str(p))
    chunks = parse_xlsx_file(p, max_chars=800, overlap=0)
    assert len(chunks) == 1
    assert chunks[0].sheet_name == "Sensors"
    assert "level_1 | 3.42" in chunks[0].content


def test_pptx(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Q3 Results"
    slide.placeholders[1].text = "Revenue grew 12 percent"
    p = tmp_path / "deck.pptx"
    prs.save(str(p))
    chunks = parse_pptx_file(p, max_chars=800, overlap=0)
    assert any(c.slide_number == 1 for c in chunks)
    assert any("Revenue" in c.content for c in chunks)


def test_html(tmp_path):
    p = tmp_path / "page.html"
    p.write_text(
        "<html><head><title>Report</title></head><body><script>evil()</script>"
        "<h1>Heading</h1><p>Main content here</p></body></html>",
        encoding="utf-8",
    )
    chunks = parse_html_file(p, max_chars=500, overlap=0)
    text = "\n".join(c.content for c in chunks)
    assert "Report" in text
    assert "evil()" not in text
